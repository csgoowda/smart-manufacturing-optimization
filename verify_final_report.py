import docx
import re
from pathlib import Path

def verify():
    filename = "Smart_Manufacturing_Optimization_Final_Report.docx"
    print(f"=== VERIFYING {filename} ===")
    doc = docx.Document(filename)
    
    forbidden = [
        "Optimal Task Planner",
        "Multiple Sclerosis",
        "Neurologist",
        "TensorFlow",
        "Keras",
        "CUDA",
        "cuDNN",
        "React",
        "Vite",
        "Turkish",
        "bilingual",
        "Pedantic Contributors",
        "Pedantic"
    ]
    
    errors = []
    
    # Check paragraphs
    for i, p in enumerate(doc.paragraphs):
        for f in forbidden:
            # check word boundary
            pattern = r'\b' + re.escape(f) + r'\b'
            if re.search(pattern, p.text, re.IGNORECASE):
                # check if it's pydantic vs pedantic
                if f.lower() == "pedantic" and "pydantic" in p.text.lower():
                    # check if the word "pedantic" specifically matched
                    if re.search(r'\bpedantic\b', p.text, re.IGNORECASE):
                        errors.append(f"Paragraph {i} matches '{f}': {p.text[:100]}")
                else:
                    errors.append(f"Paragraph {i} matches '{f}': {p.text[:100]}")
                    
    # Check tables
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for f in forbidden:
                    pattern = r'\b' + re.escape(f) + r'\b'
                    if re.search(pattern, cell.text, re.IGNORECASE):
                        if f.lower() == "pedantic" and re.search(r'\bpedantic\b', cell.text, re.IGNORECASE):
                            errors.append(f"Table {ti} R{ri}C{ci} matches '{f}': {cell.text[:80]}")
                        elif f.lower() != "pedantic":
                            errors.append(f"Table {ti} R{ri}C{ci} matches '{f}': {cell.text[:80]}")

    # Check headers and footers
    for si, s in enumerate(doc.sections):
        for p in s.header.paragraphs:
            for f in forbidden:
                if re.search(r'\b' + re.escape(f) + r'\b', p.text, re.IGNORECASE):
                    errors.append(f"Sec {si} Header matches '{f}': {p.text}")
        for p in s.footer.paragraphs:
            for f in forbidden:
                if re.search(r'\b' + re.escape(f) + r'\b', p.text, re.IGNORECASE):
                    errors.append(f"Sec {si} Footer matches '{f}': {p.text}")
                    
    if errors:
        print(f"FAILED: Found {len(errors)} forbidden string occurrences:")
        for e in errors:
            print("  -", e)
    else:
        print("PASSED: Zero occurrences of forbidden / outdated strings found!")
        
    print(f"\nDocument Statistics:")
    print(f"  Paragraphs: {len(doc.paragraphs)}")
    print(f"  Tables: {len(doc.tables)}")
    print(f"  Sections: {len(doc.sections)}")
    
    # Check figures
    figs = [p.text for p in doc.paragraphs if p.text.startswith("Figure")]
    print(f"\nFigures found ({len(figs)}):")
    for f in figs:
        print("  ", f)
        
    # Check tables captions
    tbls = [p.text for p in doc.paragraphs if p.text.startswith("Table")]
    print(f"\nTables found ({len(tbls)}):")
    for tb in tbls:
        print("  ", tb)

    # Check images
    img_rels = [rel.target_ref for rel in doc.part.rels.values() if 'image' in rel.target_ref]
    print(f"\nImages embedded in document ({len(img_rels)}):")
    for im in set(img_rels):
        print("  ", im)

if __name__ == "__main__":
    verify()
