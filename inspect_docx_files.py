import docx

def inspect_format_details():
    doc = docx.Document("Report Format.docx")
    print("=== FORMAT DETAILS ===")
    for s in doc.sections:
        print(f"Margins: top={s.top_margin.inches}in, bottom={s.bottom_margin.inches}in, left={s.left_margin.inches}in, right={s.right_margin.inches}in")
        print(f"Page size: {s.page_width.inches}in x {s.page_height.inches}in")
        
    print("\n--- STYLES ---")
    for style in doc.styles:
        if style.name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'Title', 'Subtitle']:
            font = style.font
            print(f"Style: {style.name}, Font: {font.name}, Size: {font.size.pt if font.size else 'None'}, Bold: {font.bold}")

    print("\n--- SAMPLE PARAGRAPHS IN FORMAT ---")
    for i in [1, 2, 4, 5, 6, 7, 8, 9, 10, 12, 17, 18, 21, 23, 25, 28, 32, 46, 47, 49, 52, 82, 83]:
        if i < len(doc.paragraphs):
            p = doc.paragraphs[i]
            r_fonts = [f"{r.font.name} {r.font.size.pt if r.font.size else 'def'} bold={r.bold}" for r in p.runs]
            print(f"P{i} [{p.style.name}] align={p.alignment}: '{p.text[:60]}' | Runs: {r_fonts[:2]}")

if __name__ == "__main__":
    inspect_format_details()
