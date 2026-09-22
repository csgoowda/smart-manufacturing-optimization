"""
Generate Smart_Manufacturing_Optimization_Final_Report.docx.

Fully complies with JSS Science and Technology University format (Report Format.docx)
and incorporates all genuine technical content from report.docx reorganized into:
- Cover Page
- Acceptance Letter
- Certificate
- Acknowledgement
- Abstract
- Table of Contents
- List of Figures
- List of Tables
- Chapter 1: Introduction
- Chapter 2: Organization Profile
- Chapter 3: Project Work (3.1 to 3.6)
- Chapter 4: Conclusions (4.1, 4.2)
- References (IEEE)
- Appendices (A, B, C)
"""

import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

SCRIPT_DIR = Path(__file__).parent
OUTPUT_FILE = SCRIPT_DIR / "Smart_Manufacturing_Optimization_Final_Report.docx"
IMG_DIR = SCRIPT_DIR / "extracted_images"
DOCS_DIR = SCRIPT_DIR / "docs"

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_para(doc, text="", bold=False, italic=False, size=12,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0,
             line_spacing=18, font_name="Times New Roman", color=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if line_spacing:
        p.paragraph_format.line_spacing = Pt(line_spacing)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = font_name
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    return p

def add_body(doc, text, indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(18)
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.27)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return p

def add_bullet(doc, title, text=""):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(16)
    
    run_t = p.add_run(title)
    run_t.bold = True
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(12)
    
    if text:
        run_d = p.add_run(f": {text}")
        run_d.font.name = "Times New Roman"
        run_d.font.size = Pt(12)
    return p

def add_num_item(doc, num_str, title, text=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = Pt(16)
    
    run_n = p.add_run(f"{num_str}. ")
    run_n.bold = True
    run_n.font.name = "Times New Roman"
    run_n.font.size = Pt(12)
    
    run_t = p.add_run(title)
    run_t.bold = True
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(12)
    
    if text:
        run_d = p.add_run(f": {text}")
        run_d.font.name = "Times New Roman"
        run_d.font.size = Pt(12)
    return p

def add_chapter_heading(doc, num_str, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    
    if num_str and str(num_str).strip():
        run = p.add_run(f"CHAPTER {num_str}: {title.upper()}")
    else:
        run = p.add_run(title.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_sec_heading(doc, num_str, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(f"{num_str}  {title}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_subsec_heading(doc, num_str, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(f"{num_str}  {title}")
    run.bold = True
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_fig_caption(doc, fig_num_str, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.keep_with_next = False
    
    run_lbl = p.add_run(f"Figure {fig_num_str}: ")
    run_lbl.bold = True
    run_lbl.font.name = "Times New Roman"
    run_lbl.font.size = Pt(10)
    
    run_t = p.add_run(title)
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(10)
    return p

def add_tbl_caption(doc, tbl_num_str, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run_lbl = p.add_run(f"Table {tbl_num_str}: ")
    run_lbl.bold = True
    run_lbl.font.name = "Times New Roman"
    run_lbl.font.size = Pt(10)
    
    run_t = p.add_run(title)
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(10)
    return p

def add_image_box(doc, img_path, fig_num_str, caption, width_in=5.8):
    if not os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"[Diagram / Screenshot: {caption}]")
        run.italic = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        add_fig_caption(doc, fig_num_str, caption)
        return
        
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(img_path), width=Inches(width_in))
    add_fig_caption(doc, fig_num_str, caption)

def build_table(doc, headers, data_rows, col_widths=None, header_bg="EAECEE"):
    t = doc.add_table(rows=1 + len(data_rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    
    # Header Row
    hdr_row = t.rows[0]
    for j, h in enumerate(headers):
        cell = hdr_row.cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(13)
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        set_cell_shading(cell, header_bg)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        
    # Data Rows
    for i, row in enumerate(data_rows):
        r = t.rows[i + 1]
        bg = "FFFFFF" if i % 2 == 0 else "F9FBFD"
        for j, val in enumerate(row):
            cell = r.cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j > 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = Pt(13)
            run = p.add_run(str(val))
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
                
    # Add spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    return t

print("Helper functions defined.")
